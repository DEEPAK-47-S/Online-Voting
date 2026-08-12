from flask import Flask, render_template, request, redirect, url_for, session, flash, jsonify
from werkzeug.utils import secure_filename
import os
import sqlite3
from functools import wraps
from datetime import datetime
import csv
import io

app = Flask(__name__)
app.secret_key = "CHANGE_THIS_SECRET_KEY"
DB = "voting.db"
UPLOAD_FOLDER = os.path.join(app.root_path, "static", "uploads")
ALLOWED_EXTENSIONS = {"png", "jpg", "jpeg", "webp"}
ALLOWED_VOTER_IMPORT_EXTENSIONS = {"csv", "xlsx", "xls"}
os.makedirs(UPLOAD_FOLDER, exist_ok=True)

ADMIN_USERNAME = "admin"
ADMIN_PASSWORD = "admin123"   # Change before deployment


def get_db():
    conn = sqlite3.connect(DB, timeout=20)
    conn.row_factory = sqlite3.Row
    return conn

def allowed_file(filename):
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS

def save_candidate_photo(file):
    if not file or not file.filename:
        return ""
    if not allowed_file(file.filename):
        raise ValueError("Only PNG, JPG, JPEG and WEBP images are allowed.")
    import uuid
    filename = secure_filename(file.filename)
    filename = f"{uuid.uuid4().hex}_{filename}"
    file.save(os.path.join(UPLOAD_FOLDER, filename))
    return "/static/uploads/" + filename



def init_db():
    conn = get_db()
    conn.executescript("""
    CREATE TABLE IF NOT EXISTS candidates (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        name TEXT NOT NULL,
        department TEXT NOT NULL,
        year TEXT NOT NULL,
        symbol TEXT,
        manifesto TEXT,
        photo TEXT,
        votes INTEGER NOT NULL DEFAULT 0
    );

    CREATE TABLE IF NOT EXISTS voters (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        voter_id TEXT UNIQUE NOT NULL,
        name TEXT NOT NULL,
        department TEXT,
        has_voted INTEGER NOT NULL DEFAULT 0
    );

    CREATE TABLE IF NOT EXISTS votes (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        voter_id TEXT UNIQUE NOT NULL,
        candidate_id INTEGER NOT NULL,
        voted_at TEXT NOT NULL,
        FOREIGN KEY(candidate_id) REFERENCES candidates(id)
    );
    """)
    conn.commit()
    conn.close()


def admin_required(fn):
    @wraps(fn)
    def wrapper(*args, **kwargs):
        if not session.get("admin"):
            return redirect(url_for("admin_login"))
        return fn(*args, **kwargs)
    return wrapper


@app.route("/")
def index():
    return render_template("index.html")


@app.route("/vote", methods=["GET", "POST"])
def vote():
    conn = get_db()

    if request.method == "POST":
        voter_id = request.form.get("voter_id", "").strip().upper()
        candidate_id = request.form.get("candidate_id")

        voter = conn.execute(
            "SELECT * FROM voters WHERE voter_id = ?", (voter_id,)
        ).fetchone()

        if not voter:
            conn.close()
            return render_template("vote.html", error="Invalid Voter ID.", candidates=[])

        if voter["has_voted"]:
            conn.close()
            return render_template(
                "vote.html",
                error="This Voter ID has already submitted a vote.",
                candidates=[]
            )

        candidate = conn.execute(
            "SELECT * FROM candidates WHERE id = ?", (candidate_id,)
        ).fetchone()

        if not candidate:
            conn.close()
            return render_template("vote.html", error="Please select a valid candidate.", candidates=[])

        try:
            conn.execute(
                "INSERT INTO votes (voter_id, candidate_id, voted_at) VALUES (?, ?, ?)",
                (voter_id, candidate_id, datetime.now().isoformat(timespec="seconds"))
            )
            conn.execute(
                "UPDATE candidates SET votes = votes + 1 WHERE id = ?",
                (candidate_id,)
            )
            conn.execute(
                "UPDATE voters SET has_voted = 1 WHERE voter_id = ?",
                (voter_id,)
            )
            conn.commit()
        except sqlite3.IntegrityError:
            conn.rollback()
            conn.close()
            return render_template(
                "vote.html",
                error="Vote could not be submitted. This voter may have already voted.",
                candidates=[]
            )

        voter_name = voter["name"]
        voter_department = voter["department"] or ""
        conn.close()
        return render_template(
            "success.html",
            voter_id=voter_id,
            voter_name=voter_name,
            voter_department=voter_department
        )

    candidates = conn.execute(
        "SELECT * FROM candidates ORDER BY id"
    ).fetchall()
    conn.close()
    return render_template("vote.html", candidates=candidates)


@app.route("/admin/login", methods=["GET", "POST"])
def admin_login():
    if request.method == "POST":
        username = request.form.get("username", "")
        password = request.form.get("password", "")

        if username == ADMIN_USERNAME and password == ADMIN_PASSWORD:
            session["admin"] = True
            return redirect(url_for("dashboard"))

        return render_template("admin_login.html", error="Invalid username or password.")

    return render_template("admin_login.html")


@app.route("/admin/logout")
def admin_logout():
    session.clear()
    return redirect(url_for("admin_login"))


@app.route("/admin")
@admin_required
def dashboard():
    conn = get_db()
    candidates = conn.execute("SELECT * FROM candidates ORDER BY id").fetchall()
    voter_count = conn.execute("SELECT COUNT(*) AS c FROM voters").fetchone()["c"]
    voted_count = conn.execute("SELECT COUNT(*) AS c FROM voters WHERE has_voted=1").fetchone()["c"]
    total_votes = conn.execute("SELECT COUNT(*) AS c FROM votes").fetchone()["c"]
    voters = conn.execute("SELECT voter_id, name, department, has_voted FROM voters ORDER BY id DESC").fetchall()
    conn.close()

    return render_template(
        "dashboard.html",
        candidates=candidates,
        voters=voters,
        voter_count=voter_count,
        voted_count=voted_count,
        total_votes=total_votes
    )


@app.route("/admin/candidate/add", methods=["POST"])
@admin_required
def add_candidate():
    name = request.form.get("name", "").strip()
    department = request.form.get("department", "").strip()
    year = request.form.get("year", "").strip()
    symbol = request.form.get("symbol", "").strip()
    manifesto = request.form.get("manifesto", "").strip()
    photo_file = request.files.get("photo")

    if not name or not department or not year:
        flash("Name, department and year are required.")
        return redirect(url_for("dashboard"))

    conn = get_db()
    count = conn.execute("SELECT COUNT(*) AS c FROM candidates").fetchone()["c"]
    if count >= 11:
        conn.close()
        flash("Maximum 11 candidates are allowed.")
        return redirect(url_for("dashboard"))

    try:
        photo = save_candidate_photo(photo_file)
    except ValueError as e:
        conn.close()
        flash(str(e))
        return redirect(url_for("dashboard"))

    conn.execute("""
        INSERT INTO candidates(name, department, year, symbol, manifesto, photo)
        VALUES (?, ?, ?, ?, ?, ?)
    """, (name, department, year, symbol, manifesto, photo))
    conn.commit()
    conn.close()
    flash("Candidate added successfully.")
    return redirect(url_for("dashboard"))


@app.route("/admin/candidate/delete/<int:candidate_id>", methods=["POST"])
@admin_required
def delete_candidate(candidate_id):
    conn = get_db()
    candidate = conn.execute(
        "SELECT votes FROM candidates WHERE id=?", (candidate_id,)
    ).fetchone()

    if candidate and candidate["votes"] > 0:
        flash("A candidate with votes cannot be deleted.")
    else:
        conn.execute("DELETE FROM candidates WHERE id=?", (candidate_id,))
        conn.commit()
        flash("Candidate deleted.")
    conn.close()
    return redirect(url_for("dashboard"))


@app.route("/admin/voter/add", methods=["POST"])
@admin_required
def add_voter():
    voter_id = request.form.get("voter_id", "").strip().upper()
    name = request.form.get("name", "").strip()
    department = request.form.get("department", "").strip()

    if not voter_id or not name:
        flash("Voter ID and name are required.")
        return redirect(url_for("dashboard"))

    conn = get_db()
    try:
        conn.execute(
            "INSERT INTO voters(voter_id, name, department) VALUES (?, ?, ?)",
            (voter_id, name, department)
        )
        conn.commit()
        flash("Voter added successfully.")
    except sqlite3.IntegrityError:
        flash("Voter ID already exists.")
    conn.close()
    return redirect(url_for("dashboard"))


@app.route("/admin/voter/bulk-upload", methods=["POST"])
@admin_required
def bulk_upload_voters():
    file = request.files.get("voter_file")
    if not file or not file.filename:
        flash("Please choose a CSV or Excel file.")
        return redirect(url_for("dashboard"))

    ext = file.filename.rsplit(".", 1)[-1].lower() if "." in file.filename else ""
    if ext not in ALLOWED_VOTER_IMPORT_EXTENSIONS:
        flash("Only CSV, XLSX or XLS voter files are supported.")
        return redirect(url_for("dashboard"))

    try:
        rows = []
        if ext == "csv":
            text = file.read().decode("utf-8-sig")
            reader = csv.DictReader(io.StringIO(text))
            if not reader.fieldnames:
                raise ValueError("The CSV file has no header row.")
            headers = {h.strip().lower().replace(" ", "_") for h in reader.fieldnames if h}
            if "voter_id" not in headers or "name" not in headers:
                raise ValueError("CSV must contain voter_id and name columns. department is optional.")
            for row in reader:
                normalized = {(k or "").strip().lower().replace(" ", "_"): (v or "").strip() for k, v in row.items()}
                rows.append((normalized.get("voter_id", ""), normalized.get("name", ""), normalized.get("department", "")))
        else:
            from openpyxl import load_workbook
            wb = load_workbook(file, read_only=True, data_only=True)
            ws = wb.active
            data = list(ws.iter_rows(values_only=True))
            if not data:
                raise ValueError("The Excel file is empty.")
            header_row = [str(x).strip().lower().replace(" ", "_") if x is not None else "" for x in data[0]]
            if "voter_id" not in header_row or "name" not in header_row:
                raise ValueError("Excel must contain voter_id and name columns. department is optional.")
            id_i = header_row.index("voter_id")
            name_i = header_row.index("name")
            dept_i = header_row.index("department") if "department" in header_row else None
            for values in data[1:]:
                get = lambda i: "" if i is None or i >= len(values) or values[i] is None else str(values[i]).strip()
                rows.append((get(id_i), get(name_i), get(dept_i)))
            wb.close()

        conn = get_db()
        inserted = 0
        skipped = 0
        invalid = 0
        seen = set()
        for voter_id, name, department in rows:
            voter_id = voter_id.upper().strip()
            if not voter_id or not name or voter_id in seen:
                invalid += 1
                continue
            seen.add(voter_id)
            try:
                conn.execute(
                    "INSERT INTO voters(voter_id, name, department) VALUES (?, ?, ?)",
                    (voter_id, name, department)
                )
                inserted += 1
            except sqlite3.IntegrityError:
                skipped += 1
        conn.commit()
        conn.close()
        flash(f"Bulk upload complete: {inserted} added, {skipped} already registered, {invalid} invalid/duplicate rows skipped.")
    except UnicodeDecodeError:
        flash("CSV must be saved as UTF-8. Please use the provided template.")
    except Exception as e:
        flash(f"Bulk upload failed: {e}")
    return redirect(url_for("dashboard"))


@app.route("/admin/voter/template")
@admin_required
def voter_template():
    import csv, io
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["voter_id", "name", "department"])
    writer.writerow(["NIET2026001", "Sample Student", "CSE"])
    writer.writerow(["NIET2026002", "Sample Student 2", "ECE"])
    from flask import Response
    return Response(output.getvalue(), mimetype="text/csv", headers={
        "Content-Disposition": "attachment; filename=voter_upload_template.csv"
    })


@app.route("/admin/results")
@admin_required
def results():
    conn = get_db()
    results = conn.execute("""
        SELECT id, name, department, year, symbol, votes
        FROM candidates
        ORDER BY votes DESC, name ASC
    """).fetchall()
    conn.close()
    return render_template("results.html", results=results)


@app.route("/api/admin/results")
@admin_required
def api_admin_results():
    conn = get_db()
    candidates = conn.execute("""
        SELECT id, name, department, year, symbol, photo, votes
        FROM candidates ORDER BY votes DESC, name ASC
    """).fetchall()
    voter_count = conn.execute("SELECT COUNT(*) AS c FROM voters").fetchone()["c"]
    voted_count = conn.execute("SELECT COUNT(*) AS c FROM voters WHERE has_voted=1").fetchone()["c"]
    total_votes = conn.execute("SELECT COUNT(*) AS c FROM votes").fetchone()["c"]
    conn.close()
    percentage = round(voted_count / voter_count * 100, 1) if voter_count else 0
    return jsonify({
        "total_votes": total_votes,
        "voted_count": voted_count,
        "remaining": voter_count - voted_count,
        "percentage": percentage,
        "candidates": [dict(c) for c in candidates]
    })


@app.route("/admin/reset-votes", methods=["POST"])
@admin_required
def reset_votes():
    # Useful for testing before the real election.
    conn = get_db()
    conn.execute("DELETE FROM votes")
    conn.execute("UPDATE candidates SET votes=0")
    conn.execute("UPDATE voters SET has_voted=0")
    conn.commit()
    conn.close()
    flash("All votes have been reset.")
    return redirect(url_for("dashboard"))


@app.route("/admin/export")
@admin_required
def export_data():
    conn = get_db()
    rows = conn.execute("""
        SELECT v.voter_id, v.name AS voter_name, v.department,
               c.name AS candidate_name, c.symbol, vt.voted_at
        FROM votes vt
        JOIN voters v ON v.voter_id = vt.voter_id
        JOIN candidates c ON c.id = vt.candidate_id
        ORDER BY vt.voted_at
    """).fetchall()
    conn.close()

    import csv, io
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["Voter ID", "Voter Name", "Department", "Candidate", "Symbol", "Voted At"])
    for row in rows:
        writer.writerow(list(row))

    from flask import Response
    return Response(
        output.getvalue(),
        mimetype="text/csv",
        headers={"Content-Disposition": "attachment; filename=election_votes.csv"}
    )


if __name__ == "__main__":
    init_db()
    app.run(host="0.0.0.0", port=5000, debug=True)
